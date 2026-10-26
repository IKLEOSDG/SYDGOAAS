from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import re
D=Document();sec=D.sections[0];sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=sec.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2']:
 s=D.styles[name];s.font.name='Microsoft YaHei';s._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');s.font.color.rgb=RGBColor(0,0,0)
D.styles['Normal'].font.size=Pt(10.5);D.styles['Normal'].paragraph_format.space_after=Pt(7)
D.styles['Title'].font.size=Pt(24)
for name in ['Heading 1','Heading 2']:D.styles[name].font.size=Pt(16 if name=='Heading 1' else 12)
def addmd(text):
 lines=text.splitlines();i=0
 while i<len(lines):
  l=lines[i].strip()
  if l.startswith('|'):
   block=[]
   while i<len(lines) and lines[i].strip().startswith('|'):
    row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
    if not all(re.fullmatch(r'[-: ]+',x) for x in row):block.append(row)
    i+=1
   t=D.add_table(rows=0,cols=len(block[0]));t.autofit=False
   for ri,row in enumerate(block):
    cells=t.add_row().cells
    for c,txt in zip(cells,row):
     c.text=txt
     tcpr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
     for edge in ['top','left','bottom','right']:
      e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
     tcpr.append(b)
     if ri==0:
      sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EFF7');tcpr.append(sh)
     for p in c.paragraphs:
      p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
      for r in p.runs:r.font.size=Pt(9);r.bold=ri==0
    if ri==0:
     el=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(el)
   D.add_paragraph();continue
  if l.startswith('# '):D.add_paragraph(l[2:],'Title')
  elif l.startswith('## '):D.add_heading(re.sub(r'^[0-9]+\.\s*','',l[3:]),level=1)
  elif l:D.add_paragraph(l)
  i+=1
addmd(Path('PRD.md').read_text(encoding='utf-8-sig'))
D.add_heading('移动端交互要求',level=1)
D.add_paragraph('手机端采用抽屉式主导航，十三册导航可横向滑动；表单单列排列，表格在独立区域横向滚动。弹窗使用底部面板，主要操作触控区域不少于44像素。导出字段以紧凑选项卡排列，复选框与文字保持关联。')
D.add_page_break();D.add_heading('附录 原始模板字段映射',level=1)
D.add_paragraph('以下按实际模板提取业务字段。主键、操作人、操作时间及员工关联字段由正式系统处理。星号沿用原模板，不代表所有情形均需无条件填写。')
text=Path('字段映射.md').read_text(encoding='utf-8-sig');addmd('\n'.join(text.splitlines()[4:]))
D.save('沈阳数据集团人员档案平台PRD.docx')
